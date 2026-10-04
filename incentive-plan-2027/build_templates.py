# -*- coding: utf-8 -*-
"""Build the bilingual Google Sheets template workbook for the 2027 incentive plan.
Output: incentive_plan_2027_templates.xlsx (upload to Google Drive -> converts to Sheets)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")
TITLE_FONT = Font(name=FONT, bold=True, size=13)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="595959")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 42

def body(ws, row, values, example=False):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = INPUT_FONT if example else BODY_FONT
        c.alignment = WRAP; c.border = BORDER
        if example:
            c.fill = YELLOW

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub; ws["A2"].font = NOTE_FONT
    ws.row_dimensions[1].height = 22

def dv_list(ws, options, cell_range):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "请从下拉列表选择 / Vui lòng chọn từ danh sách"
    ws.add_data_validation(dv); dv.add(cell_range)

DEPTS = ["生产 Sản xuất", "质量 Chất lượng", "人事 Nhân sự", "进出口 XNK", "财务 Tài chính",
         "采购 Mua hàng", "仓库 Kho", "班组 Tổ sản xuất"]
SECTIONS = ["1 本月完成 Đã hoàn thành", "2 进行中 Đang thực hiện", "3 问题与预警 Vấn đề & cảnh báo",
            "4 改善建议 Đề xuất cải tiến", "5 下月计划 Kế hoạch tháng sau"]
LEDGER_TYPES = ["改善目标 Mục tiêu cải tiến", "关键节点 Mốc quan trọng", "紧急处理 Xử lý khẩn cấp"]
STATUS = ["未开始 Chưa bắt đầu", "进行中 Đang thực hiện", "已完成 Hoàn thành", "延期 Trễ hạn", "已关闭 Đã đóng"]
YN = ["是 Có", "否 Không"]
PROP_TYPE = ["主动提案 Chủ động", "指派任务 Được giao"]
PROP_SCOPE = ["部门内 Trong bộ phận", "跨部门 Liên bộ phận"]
PROP_STATUS = ["提案 Đề xuất", "评审 Thẩm định", "执行 Thực hiện", "验证 Xác minh", "结案 Kết thúc", "否决 Từ chối"]
FIN = ["是 Có", "否 Không", "不需要 Không cần"]
CONTRIB_TYPE = ["生产交付 Giao hàng", "人员安全 An toàn", "团队管理 Quản lý đội ngũ"]
ENV = ["政策 Chính sách", "技术 Kỹ thuật", "资金 Tài chính", "理念 Tư duy"]
GM = ["待定 Chờ xét", "认定 Công nhận", "不认定 Không công nhận"]

# ---------------- 0 说明 ----------------
ws = wb.active; ws.title = "说明 Hướng dẫn"
title(ws, "2027年度专项激励计划 · 数据模板使用说明 / Hướng dẫn sử dụng mẫu dữ liệu")
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 62; ws.column_dimensions["C"].width = 62
rows = [
    ("", "中文", "Tiếng Việt"),
    ("1", "黄色底、蓝色字的行是示例，看懂格式后可以删除。", "Dòng nền vàng chữ xanh là ví dụ, hiểu định dạng rồi có thể xóa."),
    ("2", "每个部门复制本文件，改名为“部门名_2027”，放在部门文件夹内。只填本部门的数据。", "Mỗi bộ phận sao chép tệp này, đổi tên “Tên bộ phận_2027”, đặt trong thư mục bộ phận. Chỉ điền dữ liệu của bộ phận mình."),
    ("3", "“月度工作单”每月 5 日前填好上个月的内容，一件事一行。五个栏目都要有内容，没有就写“无”。", "“Phiếu công việc tháng” điền nội dung tháng trước, trước ngày 5 hàng tháng, mỗi việc một dòng. Cả 5 mục đều phải có nội dung, không có thì ghi “Không”."),
    ("4", "“部门台账”有变化就更新“最近更新”和“进展说明”。进行中的事项超过 30 天不更新会被视为停滞。", "“Sổ theo dõi” có thay đổi thì cập nhật “Cập nhật gần nhất” và “Tiến độ”. Việc đang làm quá 30 ngày không cập nhật sẽ bị coi là đình trệ."),
    ("5", "“降本增效提案”一个提案一行，证据文件放在提案文件夹，链接填到最后一列。节省金额按年化计算。", "“Đề xuất cải tiến” mỗi đề xuất một dòng, hồ sơ chứng minh để trong thư mục đề xuất, dán liên kết vào cột cuối. Số tiền tiết kiệm tính theo năm."),
    ("6", "“关键贡献登记”只登记有事前预警记录的事件，预警日期必须早于问题发生日期。", "“Đóng góp quan trọng” chỉ ghi các sự việc có bản ghi cảnh báo trước; ngày cảnh báo phải sớm hơn ngày sự việc xảy ra."),
    ("7", "编号规则：部门代码-年份-三位序号，例如 SX-2027-001。部门代码：SX 生产、CL 质量、NS 人事、XNK 进出口、TC 财务、MH 采购、KHO 仓库、TO 班组。", "Quy tắc mã: Mã bộ phận-Năm-Số thứ tự 3 chữ số, ví dụ SX-2027-001. Mã bộ phận: SX Sản xuất, CL Chất lượng, NS Nhân sự, XNK Xuất nhập khẩu, TC Tài chính, MH Mua hàng, KHO Kho, TO Tổ sản xuất."),
    ("8", "AI 每季度读取这些表格和部门文件夹，给出等级建议和证据清单，总经理最终确定。没有记录的工作无法被评定。", "AI đọc các bảng này và thư mục bộ phận mỗi quý, đưa ra xếp loại đề xuất kèm danh mục bằng chứng; Tổng Giám đốc quyết định cuối cùng. Công việc không có bản ghi thì không thể đánh giá."),
    ("9", "“参数”页由总经理维护，其他人不要修改。", "Trang “Tham số” do Tổng Giám đốc quản lý, người khác không chỉnh sửa."),
]
for r, vals in enumerate(rows, start=4):
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v); c.alignment = WRAP; c.border = BORDER
        c.font = HDR_FONT if r == 4 else BODY_FONT
        if r == 4: c.fill = HDR_FILL

# ---------------- 1 月度工作单 ----------------
ws = wb.create_sheet("月度工作单 Phiếu tháng")
title(ws, "月度工作单 / Phiếu công việc tháng", "一件事一行；每月 5 日前填完上月 / Mỗi việc một dòng; hoàn thành trước ngày 5 tháng sau")
H = ["月份\nTháng (yyyy-mm)", "部门\nBộ phận", "填写人\nNgười điền", "栏目\nMục", "内容\nNội dung",
     "数据或结果\nSố liệu / kết quả", "涉及部门\nBộ phận liên quan", "文件链接\nLiên kết tệp", "填写日期\nNgày điền"]
W = [12, 16, 14, 24, 48, 28, 18, 30, 13]
header(ws, 4, H, W)
ex = [
    ["'2027-01", DEPTS[0], "Nguyễn Văn A", SECTIONS[0], "完成 1 号机换模时间从 90 分钟缩短到 60 分钟 / Rút ngắn thời gian đổi khuôn máy số 1 từ 90 xuống 60 phút", "换模记录 12 次，平均 58 分钟 / 12 lần, TB 58 phút", "质量 Chất lượng", "https://drive.google.com/...", "2027-02-03"],
    ["'2027-01", DEPTS[0], "Nguyễn Văn A", SECTIONS[2], "2 号机牵引轮磨损，预计 3 月前需更换，否则影响 3 月订单交付 / Bánh kéo máy 2 mòn, cần thay trước tháng 3, nếu không ảnh hưởng giao hàng tháng 3", "磨损量 1.2 mm / Độ mòn 1,2 mm", "采购 Mua hàng", "", "2027-02-03"],
]
for i, row in enumerate(ex, start=5): body(ws, i, row, example=True)
dv_list(ws, DEPTS, "B5:B400"); dv_list(ws, SECTIONS, "D5:D400"); dv_list(ws, DEPTS, "G5:G400")
ws.freeze_panes = "A5"

# ---------------- 2 部门台账 ----------------
ws = wb.create_sheet("部门台账 Sổ theo dõi")
title(ws, "部门年度项目管理台账 / Sổ theo dõi công việc bộ phận", "管理层可协同编辑、新增 / Cấp quản lý có thể cùng chỉnh sửa, bổ sung")
H = ["编号\nMã", "类型\nLoại", "部门\nBộ phận", "事项\nNội dung công việc", "负责人\nNgười phụ trách",
     "登记日期\nNgày ghi", "截止日期\nHạn hoàn thành", "状态\nTrạng thái", "最近更新\nCập nhật gần nhất",
     "进展说明\nTiến độ", "关联提案编号\nMã đề xuất liên quan", "是否逾期\nQuá hạn", "停滞天数\nSố ngày chưa cập nhật"]
W = [14, 20, 16, 40, 14, 12, 12, 18, 12, 36, 16, 10, 12]
header(ws, 4, H, W)
ex = [
    ["SX-2027-001", LEDGER_TYPES[0], DEPTS[0], "降低 1 号机开机废料 / Giảm phế liệu khởi động máy 1", "Nguyễn Văn A", "2027-01-10", "2027-03-31", STATUS[1], "2027-02-03", "已收集 1 月数据，2 月试行新参数 / Đã thu số liệu tháng 1, tháng 2 thử tham số mới", "SX-2027-001", None, None],
    ["KHO-2027-001", LEDGER_TYPES[2], DEPTS[6], "旺季仓位不足，需在 11 月 15 日前腾出 B 区 / Thiếu chỗ kho mùa cao điểm, cần giải phóng khu B trước 15/11", "Trần Thị B", "2026-11-02", "2026-11-15", STATUS[2], "2026-11-14", "已完成，腾出 120 m² / Đã xong, giải phóng 120 m²", "", None, None],
]
for i, row in enumerate(ex, start=5): body(ws, i, row, example=True)
for r in range(5, 30):
    ws.cell(row=r, column=12, value=f'=IF(AND(G{r}<>"",H{r}<>"已完成 Hoàn thành",H{r}<>"已关闭 Đã đóng",G{r}<TODAY()),"逾期 Quá hạn","")').font = BODY_FONT
    ws.cell(row=r, column=13, value=f'=IF(AND(I{r}<>"",H{r}="进行中 Đang thực hiện"),TODAY()-I{r},"")').font = BODY_FONT
dv_list(ws, LEDGER_TYPES, "B5:B300"); dv_list(ws, DEPTS, "C5:C300"); dv_list(ws, STATUS, "H5:H300")
ws.freeze_panes = "A5"

# ---------------- 3 降本增效提案 ----------------
ws = wb.create_sheet("提案 Đề xuất")
title(ws, "降本增效提案登记 / Đăng ký đề xuất giảm chi phí, tăng hiệu quả", "分成比例、封顶、财务确认门槛来自“参数”页 / Tỷ lệ chia thưởng, mức tối đa, ngưỡng xác nhận tài chính lấy từ trang “Tham số”")
H = ["提案编号\nMã đề xuất", "部门\nBộ phận", "提案人\nNgười đề xuất", "类型\nLoại", "范围\nPhạm vi", "提案日期\nNgày đề xuất",
     "提案名称\nTên đề xuất", "现状与问题\nHiện trạng & vấn đề", "改善方案\nGiải pháp", "预计年化节省 (VND)\nTiết kiệm dự kiến/năm",
     "验证期起\nXác minh từ", "验证期止\nXác minh đến", "验证后年化节省 (VND)\nTiết kiệm đã xác minh/năm", "需财务确认\nCần tài chính xác nhận",
     "财务已确认\nTài chính đã xác nhận", "状态\nTrạng thái", "预计分成 (VND)\nTiền thưởng dự kiến", "证据文件夹链接\nLiên kết thư mục hồ sơ"]
W = [14, 16, 14, 16, 16, 12, 30, 34, 34, 16, 12, 12, 16, 12, 12, 14, 16, 30]
header(ws, 4, H, W)
ex = [["SX-2027-001", DEPTS[0], "Nguyễn Văn A", PROP_TYPE[0], PROP_SCOPE[1], "2027-01-10",
       "降低 1 号机开机废料 / Giảm phế liệu khởi động máy 1",
       "每次开机废料约 25 kg，月开机 20 次 / Mỗi lần khởi động ~25 kg phế, 20 lần/tháng",
       "调整预热程序并与质量部确定首检放行标准 / Điều chỉnh gia nhiệt và thống nhất tiêu chuẩn kiểm tra đầu với QC",
       120000000, "2027-03-01", "2027-05-31", 96000000, None, FIN[0], PROP_STATUS[4], None, "https://drive.google.com/..."]]
for i, row in enumerate(ex, start=5): body(ws, i, row, example=True)
for r in range(5, 30):
    ws.cell(row=r, column=14, value=f'=IF(J{r}="","",IF(J{r}>=参数!$B$9,"是 Có","否 Không"))').font = BODY_FONT
    ws.cell(row=r, column=17, value=f'=IF(OR(M{r}="",P{r}<>"结案 Kết thúc"),"",IF(D{r}="主动提案 Chủ động",MIN(M{r}*参数!$B$7,参数!$B$8),0))').font = BODY_FONT
    for cidx in (10, 13, 17):
        ws.cell(row=r, column=cidx).number_format = '#,##0;(#,##0);"-"'
dv_list(ws, DEPTS, "B5:B200"); dv_list(ws, PROP_TYPE, "D5:D200"); dv_list(ws, PROP_SCOPE, "E5:E200")
dv_list(ws, FIN, "O5:O200"); dv_list(ws, PROP_STATUS, "P5:P200")
ws.freeze_panes = "B5"

# ---------------- 4 关键贡献登记 ----------------
ws = wb.create_sheet("关键贡献 Đóng góp")
title(ws, "关键贡献事件登记 / Đăng ký sự việc đóng góp quan trọng", "只登记有事前预警记录的事件 / Chỉ ghi sự việc có bản ghi cảnh báo trước")
H = ["编号\nMã", "部门\nBộ phận", "人员\nNgười", "类型\nLoại", "风险来源\nNguồn rủi ro", "预警日期\nNgày cảnh báo",
     "预警记录链接\nLiên kết bản ghi cảnh báo", "问题发生或预计发生日期\nNgày xảy ra / dự kiến", "事件描述\nMô tả sự việc",
     "采取的行动\nHành động đã làm", "避免的损失估计 (VND)\nThiệt hại tránh được (ước tính)", "结果\nKết quả", "总经理认定\nTGĐ xét duyệt", "预警是否早于发生\nCảnh báo trước khi xảy ra"]
W = [14, 16, 14, 18, 16, 12, 30, 14, 36, 36, 16, 24, 16, 12]
header(ws, 4, H, W)
ex = [["CL-2027-001", DEPTS[1], "Lê Văn C", CONTRIB_TYPE[0], ENV[1], "2027-02-03", "https://drive.google.com/...", "2027-03-10",
       "发现 2 号机牵引轮磨损，若不更换 3 月订单将停机 / Phát hiện bánh kéo máy 2 mòn, không thay thì tháng 3 dừng máy",
       "提前申购备件，3 月 5 日更换 / Mua phụ tùng trước, thay ngày 5/3", 80000000, "3 月订单按期交付 / Đơn tháng 3 giao đúng hạn", GM[0], None]]
for i, row in enumerate(ex, start=5): body(ws, i, row, example=True)
for r in range(5, 30):
    ws.cell(row=r, column=14, value=f'=IF(OR(F{r}="",H{r}=""),"",IF(F{r}<H{r},"是 Có","否 Không"))').font = BODY_FONT
    ws.cell(row=r, column=11).number_format = '#,##0;(#,##0);"-"'
dv_list(ws, DEPTS, "B5:B200"); dv_list(ws, CONTRIB_TYPE, "D5:D200"); dv_list(ws, ENV, "E5:E200"); dv_list(ws, GM, "M5:M200")
ws.freeze_panes = "A5"

# ---------------- 5 参数 ----------------
ws = wb.create_sheet("参数 Tham số")
title(ws, "计划参数与人员分配 / Tham số kế hoạch và phân bổ", "蓝色字为总经理设定的输入值 / Chữ xanh là giá trị đầu vào do Tổng Giám đốc thiết lập")
ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 18; ws.column_dimensions["F"].width = 18; ws.column_dimensions["G"].width = 40
params = [
    ("参数 Tham số", "数值 Giá trị", "说明 Ghi chú"),
    ("保底奖金池 (VND) / Quỹ thưởng tối thiểu", 250000000, "公司承诺全额发放 / Công ty cam kết phát đủ"),
    ("计划年度 / Năm kế hoạch", "'2027", "2027-01-01 至 2027-12-31"),
    ("降本增效分成比例 / Tỷ lệ chia thưởng tiết kiệm", 0.15, "按验证后年化节省 / Theo tiết kiệm đã xác minh (năm)"),
    ("单个提案分成上限 (VND) / Mức thưởng tối đa mỗi đề xuất", 100000000, "总额不封顶 / Tổng không giới hạn"),
    ("需财务确认的年化节省门槛 (VND) / Ngưỡng cần tài chính xác nhận", 50000000, "低于此值由总经理按等级认定 / Dưới mức này TGĐ xét theo xếp loại"),
    ("月度工作单截止日 (每月) / Hạn nộp phiếu tháng", 5, "次月 5 日前 / Trước ngày 5 tháng sau"),
    ("台账停滞天数阈值 / Ngưỡng ngày đình trệ", 30, "进行中事项超过此天数未更新 / Việc đang làm quá số ngày này chưa cập nhật"),
]
for r, vals in enumerate(params, start=4):
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v); c.border = BORDER; c.alignment = WRAP
        if r == 4: c.font = HDR_FONT; c.fill = HDR_FILL
        else: c.font = INPUT_FONT if ci == 2 else BODY_FONT
ws["B5"].number_format = '#,##0'; ws["B7"].number_format = '0%'; ws["B8"].number_format = '#,##0'; ws["B9"].number_format = '#,##0'
# NOTE: rows -> B5 pool, B7 share %, B8 cap, B9 threshold (referenced by 提案 sheet)

r0 = 14
ws.cell(row=r0 - 1, column=1, value="等级倍数 / Hệ số theo xếp loại").font = Font(name=FONT, bold=True, size=11)
grades = [("等级 Xếp loại", "倍数 Hệ số", "说明 Ghi chú"),
          ("A 优秀 Xuất sắc", 1.3, "未发完的余额分给 A 档 / Phần dư chia cho loại A"),
          ("B 达标 Đạt", 1.0, "标准额 / Mức chuẩn"),
          ("C 不足 Chưa đạt", 0.6, ""),
          ("D 未完成 Không đạt", 0.0, "")]
for i, vals in enumerate(grades):
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r0 + i, column=ci, value=v); c.border = BORDER; c.alignment = WRAP
        if i == 0: c.font = HDR_FONT; c.fill = HDR_FILL
        else: c.font = INPUT_FONT if ci == 2 else BODY_FONT

r1 = 21
ws.cell(row=r1 - 1, column=1, value="人员与标准额 / Danh sách và mức chuẩn (B 档)").font = Font(name=FONT, bold=True, size=11)
hdr = ["岗位 Vị trí", "系数 Hệ số", "姓名 Họ tên", "年度等级 Xếp loại năm", "标准额 (VND) Mức chuẩn", "应发 (VND) Thực nhận", "说明 Ghi chú"]
for ci, v in enumerate(hdr, start=1):
    c = ws.cell(row=r1, column=ci, value=v); c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BORDER; c.alignment = CENTER
people = [("生产主管 Giám sát sản xuất", 1.0), ("质量主管 Giám sát chất lượng", 1.0), ("采购主管 Trưởng mua hàng", 1.0),
          ("仓库主管 Giám sát kho", 1.0), ("人事主管 Trưởng nhân sự", 1.0), ("进出口跟单 NV chứng từ XNK", 0.7),
          ("出纳 Thủ quỹ", 0.7), ("质量助理 Trợ lý chất lượng", 0.7),
          ("班长 1 Tổ trưởng ca 1", 0.4), ("班长 2 Tổ trưởng ca 2", 0.4), ("班长 3 Tổ trưởng ca 3", 0.4), ("班长 4 Tổ trưởng ca 4", 0.4), ("班长 5 Tổ trưởng ca 5", 0.4)]
first = r1 + 1; last = r1 + len(people)
for i, (pos, k) in enumerate(people):
    r = first + i
    vals = [pos, k, "", "B 达标 Đạt",
            f"=$B$5/SUM($B${first}:$B${last})*B{r}",
            f"=IF(D{r}=\"\",\"\",E{r}*IFERROR(VLOOKUP(D{r},$A${r0+1}:$B${r0+4},2,FALSE),1))",
            ""]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v); c.border = BORDER; c.alignment = WRAP
        c.font = INPUT_FONT if ci in (2, 3, 4) else BODY_FONT
        if ci == 3: c.fill = YELLOW
    ws.cell(row=r, column=5).number_format = '#,##0'; ws.cell(row=r, column=6).number_format = '#,##0'
dv_list(ws, ["A 优秀 Xuất sắc", "B 达标 Đạt", "C 不足 Chưa đạt", "D 未完成 Không đạt"], f"D{first}:D{last}")
tr = last + 1
ws.cell(row=tr, column=1, value="合计 Tổng").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=tr, column=2, value=f"=SUM(B{first}:B{last})").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=tr, column=5, value=f"=SUM(E{first}:E{last})").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=tr, column=6, value=f"=SUM(F{first}:F{last})").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=tr, column=5).number_format = '#,##0'; ws.cell(row=tr, column=6).number_format = '#,##0'
ws.cell(row=tr + 1, column=1, value="未发余额，分给 A 档 / Phần dư chia cho loại A").font = BODY_FONT
ws.cell(row=tr + 1, column=6, value=f"=$B$5-F{tr}").number_format = '#,##0'
ws.cell(row=tr + 2, column=1, value="A 档人数 / Số người loại A").font = BODY_FONT
ws.cell(row=tr + 2, column=6, value=f'=COUNTIF(D{first}:D{last},"A 优秀 Xuất sắc")')
ws.cell(row=tr + 3, column=1, value="每位 A 档追加 / Mỗi người loại A nhận thêm").font = BODY_FONT
ws.cell(row=tr + 3, column=6, value=f"=IF(F{tr+2}=0,0,F{tr+1}/F{tr+2})").number_format = '#,##0'
for rr in range(tr, tr + 4):
    for ci in range(1, 8):
        ws.cell(row=rr, column=ci).border = BORDER

wb.save("incentive_plan_2027_templates.xlsx")
print("saved")
